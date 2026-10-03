"""
Extract every non-empty sheet from the master corpus Excel file into individual
CSV files under mathbank/src/mathbank/data/maths_corpus/.

Usage:
    python scripts/extract_corpus_csvs.py
    python scripts/extract_corpus_csvs.py --xlsx /path/to/other.xlsx
    python scripts/extract_corpus_csvs.py --out /custom/output/dir

Output:
    data/maths_corpus/<snake_case_sheet_name>.csv   — one per non-empty sheet
    data/maths_corpus/_manifest.csv                 — inventory of all exports
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from rich.console import Console
from rich.table import Table

XLSX_DEFAULT = (
    Path(__file__).resolve().parents[2]          # knowledge-bank-ingestion/
    / "archive"
    / "Tika Competition Math Prep - AMC10 AMC12 AIME MathPrize HMMT SMT PUMaC CMM CHMMC"
    / "Tika Competition Math Master Question Corpus.xlsx"
)

OUT_DEFAULT = Path(__file__).resolve().parents[1] / "src" / "mathbank" / "data" / "maths_corpus"

# Sheets that are known layout/index sheets — exported but flagged in manifest.
_LAYOUT_SHEETS = {"00 Workbook Index", "Sheet1", "Metadata Repository"}

console = Console()


def _slug(name: str) -> str:
    """Sheet name → safe snake_case filename (no extension)."""
    s = name.strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def _is_empty(rows: list[tuple]) -> bool:
    """True when the sheet has no header or only blank rows."""
    if not rows:
        return True
    header = rows[0]
    if all(v is None or str(v).strip() == "" for v in header):
        return True
    data_rows = [r for r in rows[1:] if any(v is not None and str(v).strip() for v in r)]
    return len(data_rows) == 0


def _cell_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def extract(xlsx: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    console.print(f"\n[bold]Source:[/] {xlsx}")
    console.print(f"[bold]Output:[/] {out_dir}\n")

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    manifest_rows: list[dict] = []
    table = Table("Sheet", "CSV file", "Rows", "Cols", "Status", show_lines=False)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        raw = [tuple(c for c in row) for row in ws.iter_rows(values_only=True)]

        slug = _slug(sheet_name)
        csv_name = f"{slug}.csv"
        csv_path = out_dir / csv_name

        if _is_empty(raw):
            table.add_row(sheet_name, csv_name, "—", "—", "[dim]EMPTY — skipped")
            manifest_rows.append({
                "sheet_name": sheet_name,
                "csv_file": csv_name,
                "data_rows": 0,
                "columns": 0,
                "status": "EMPTY",
                "layout_sheet": str(sheet_name in _LAYOUT_SHEETS),
            })
            continue

        # Trim trailing all-None rows.
        while raw and all(v is None for v in raw[-1]):
            raw.pop()

        # Trim trailing all-None columns.
        if raw:
            max_col = max(
                (i for row in raw for i, v in enumerate(row) if v is not None),
                default=-1,
            ) + 1
            raw = [row[:max_col] for row in raw]

        header = [_cell_str(v) or f"col_{i}" for i, v in enumerate(raw[0])]
        data_rows = raw[1:]
        n_data = len(data_rows)
        n_cols = len(header)

        with csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
            writer.writerow(header)
            for row in data_rows:
                padded = list(row) + [None] * max(0, n_cols - len(row))
                writer.writerow([_cell_str(v) for v in padded[:n_cols]])

        is_layout = sheet_name in _LAYOUT_SHEETS
        status = "LAYOUT" if is_layout else "OK"
        colour = "yellow" if is_layout else "green"
        table.add_row(sheet_name, csv_name, str(n_data), str(n_cols), f"[{colour}]{status}")
        manifest_rows.append({
            "sheet_name": sheet_name,
            "csv_file": csv_name,
            "data_rows": n_data,
            "columns": n_cols,
            "status": status,
            "layout_sheet": str(is_layout),
        })

    wb.close()

    # Write manifest.
    manifest_path = out_dir / "_manifest.csv"
    with manifest_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["sheet_name", "csv_file", "data_rows", "columns", "status", "layout_sheet"],
        )
        writer.writeheader()
        writer.writerows(manifest_rows)

    console.print(table)
    exported = sum(1 for r in manifest_rows if r["status"] in ("OK", "LAYOUT"))
    skipped = sum(1 for r in manifest_rows if r["status"] == "EMPTY")
    console.print(
        f"\n[bold green]Done.[/] {exported} CSV(s) exported, {skipped} empty sheet(s) skipped."
    )
    console.print(f"Manifest → {manifest_path}")


def main() -> None:
    p = argparse.ArgumentParser(description="Export corpus Excel sheets to CSV")
    p.add_argument("--xlsx", type=Path, default=XLSX_DEFAULT)
    p.add_argument("--out", type=Path, default=OUT_DEFAULT)
    args = p.parse_args()

    if not args.xlsx.exists():
        console.print(f"[red]Excel file not found: {args.xlsx}")
        sys.exit(1)

    extract(args.xlsx, args.out)


if __name__ == "__main__":
    main()
